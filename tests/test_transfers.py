"""Durable jobs against real PostgreSQL: interruption, replay and privacy boundaries."""

import csv
import io
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest
from openpyxl import Workbook
from test_workflow import approve
from test_workflow import make_lead as create_lead
from test_workflow import propose as proposal_for

from radar import base_queries, privacy, service, transfer_files, transfer_worker, transfers


def drain(client, ident, target="completed"):
    for _ in range(100):
        result = client.get("/api/transfers/" + str(ident)).json()
        if result["status"] == target:
            return result
        assert result["status"] in ("queued", "running"), result
        assert transfer_worker.run_one()
    pytest.fail("Job did not finish")


def upload(client, count=205, key=None, mapping=None, content=None):
    content = (
        content
        or (
            "Nome,Empresa\n"
            + "\n".join(f"Maria Exemplo,Empresa fictícia {i}" for i in range(count))
        ).encode()
    )
    path = "/api/import" if mapping else "/api/import/preview"
    import json

    response = client.post(
        path,
        files={"file": ("lista.csv", content, "text/csv")},
        data={"mapping": json.dumps(mapping)} if mapping else {},
        headers={"Idempotency-Key": str(key or uuid4())},
    )
    assert response.status_code == 202, response.text
    return response.json()


def test_upload_mapping_atomic_batches_and_lost_response(database, api_client):
    key = uuid4()
    job = upload(api_client, key=key)
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM leads").fetchone()["n"] == 0
    assert upload(api_client, key=key)["id"] == job["id"]
    preview = drain(api_client, job["id"], "awaiting_mapping")
    assert preview["total"] == 205 and len(preview["sample"]) == 5
    mapping = {"Nome": "name", "Empresa": "company"}
    assert (
        api_client.post(f"/api/transfers/{job['id']}/start", json={"mapping": mapping}).status_code
        == 202
    )
    assert transfer_worker.run_one()
    mid = api_client.get("/api/transfers/" + job["id"]).json()
    assert mid["processed"] == 100 and mid["created_count"] == 100
    # Simulate a worker dying after claiming the next batch, before committing it.
    abandoned = transfer_worker.claim()
    with database() as conn:
        conn.execute(
            "UPDATE transfer_jobs SET lease_until=now()-interval '1 second' WHERE id=%s",
            (job["id"],),
        )
    replacement = transfer_worker.claim()
    with pytest.raises(transfer_worker.LostLease):
        transfer_worker.import_batch(abandoned)
    transfer_worker.import_batch(replacement)
    done = drain(api_client, job["id"])
    assert done["processed"] == 205 and done["created_count"] == 205 and done["error_count"] == 0
    assert (
        api_client.post(f"/api/transfers/{job['id']}/start", json={"mapping": mapping}).status_code
        == 202
    )
    # File replay with a new request key is also idempotent, without requiring email.
    repeated = upload(api_client, mapping=mapping)
    result = drain(api_client, repeated["id"])
    assert result["duplicate_count"] == 205 and result["created_count"] == 0
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM leads").fetchone()["n"] == 205
        assert (
            conn.execute(
                "SELECT count(*) AS n FROM transfer_rows WHERE row_data IS NOT NULL"
            ).fetchone()["n"]
            == 0
        )
        assert conn.execute("SELECT count(*) AS n FROM transfer_blobs").fetchone()["n"] == 0
    last_page = api_client.get("/api/leads/page?page=5&page_size=50").json()
    assert len(last_page["items"]) == 5 and last_page["total"] == 205
    assert {item["company"] for item in last_page["items"]} == {
        f"Empresa fictícia {i}" for i in range(200, 205)
    }
    found = api_client.get("/api/leads/page?q=fictícia 204").json()
    assert found["total"] == 1 and found["items"][0]["company"] == "Empresa fictícia 204"


def test_replay_with_changed_upload_is_rejected(database, api_client):
    key = uuid4()
    upload(api_client, key=key)
    response = api_client.post(
        "/api/import/preview",
        files={"file": ("lista.csv", b"Empresa\nOutro")},
        headers={"Idempotency-Key": str(key)},
    )
    assert response.status_code == 409


def test_two_workers_cannot_claim_same_batch(database, api_client):
    upload(api_client)
    with ThreadPoolExecutor(max_workers=2) as pool:
        jobs = list(pool.map(lambda _: transfer_worker.claim(), range(2)))
    assert sum(job is not None for job in jobs) == 1


def test_batch_failure_rolls_back_contacts_and_progress(database, api_client, monkeypatch):
    job = upload(api_client, count=5, mapping={"Nome": "name", "Empresa": "company"})
    assert transfer_worker.run_one()
    claimed = transfer_worker.claim()
    original = service.capture
    calls = 0

    def fail_after_write(*args, **kwargs):
        nonlocal calls
        result = original(*args, **kwargs)
        calls += 1
        if calls == 3:
            raise RuntimeError("Simulated interruption inside transaction")
        return result

    with monkeypatch.context() as patch:
        patch.setattr(service, "capture", fail_after_write)
        with pytest.raises(RuntimeError):
            transfer_worker.import_batch(claimed)
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM leads").fetchone()["n"] == 0
        assert transfers.get(conn, job["id"])["processed"] == 0
    transfer_worker.import_batch(claimed)
    assert drain(api_client, job["id"])["created_count"] == 5


def test_invalid_row_and_cancellation_preserve_committed_contacts(database, api_client):
    job = upload(
        api_client,
        count=205,
        mapping={"Nome": "name", "Empresa": "company"},
        content=b"Nome,Empresa\nM. Silva,Abreviado\nMaria Silva,Valida\n",
    )
    done = drain(api_client, job["id"])
    assert (done["created_count"], done["error_count"]) == (1, 1)
    errors = api_client.get(f"/api/transfers/{job['id']}/errors").json()
    assert errors["items"][0]["line"] == 2 and "Abreviado" not in str(errors)
    second = upload(api_client, mapping={"Nome": "name", "Empresa": "company"})
    transfer_worker.run_one()
    transfer_worker.run_one()
    assert api_client.post(f"/api/transfers/{second['id']}/cancel").status_code == 200
    assert not transfer_worker.run_one()
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM leads").fetchone()["n"] == 101
        assert (
            conn.execute(
                "SELECT count(*) AS n FROM transfer_rows WHERE row_data IS NOT NULL"
            ).fetchone()["n"]
            == 0
        )


def test_csv_encoding_and_xlsx_parsing_boundaries(database, api_client, monkeypatch):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(["Nome", "Empresa"])
    for i in range(1001):
        sheet.append(["João Exemplo", f"Empresa {i}"])
    file = io.BytesIO()
    workbook.save(file)
    response = api_client.post(
        "/api/import/preview", files={"file": ("grande.xlsx", file.getvalue())}
    )
    assert response.status_code == 202
    result = drain(api_client, response.json()["id"], "awaiting_mapping")
    assert result["total"] == 1001 and result["sample"][0]["Nome"] == "João Exemplo"
    csv_job = upload(
        api_client, content="Nome;Empresa\nJoão Exemplo;Construção fictícia".encode("cp1252")
    )
    result = drain(api_client, csv_job["id"], "awaiting_mapping")
    assert result["sample"][0]["Empresa"] == "Construção fictícia"
    monkeypatch.setattr(transfer_files, "MAX_ROWS", 1)
    job = upload(api_client, count=2)
    transfer_worker.run_one()
    assert api_client.get("/api/transfers/" + job["id"]).json()["status"] == "failed"
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM leads").fetchone()["n"] == 0


def test_export_validity_and_erasure_remove_stored_copies(database, api_client):
    lead_id = create_lead(database)
    approve(
        database,
        proposal_for(database, lead_id, "website", "https://nova-empresa.invalid"),
        allow_replace=True,
    )
    response = api_client.post("/api/exports")
    assert response.status_code == 202
    done = drain(api_client, response.json()["id"])
    ready = api_client.get(f"/api/transfers/{done['id']}/download-ready")
    assert ready.status_code == 200 and ready.json()["url"] == done["download_url"]
    downloaded = api_client.get(done["download_url"])
    assert downloaded.status_code == 200
    rows = list(csv.DictReader(io.StringIO(downloaded.text.lstrip("\ufeff"))))
    assert len(rows) == 1 and rows[0]["website"] == "https://nova-empresa.invalid"
    raw = upload(api_client)
    with database() as conn:
        preview = privacy.preview(conn, lead_id)
        assert preview["counts"]["arquivos_de_listas"] == 2
        receipt = privacy.erase(conn, lead_id, preview["fingerprint"], "ELIMINAR")
        assert receipt["deleted"]
    assert api_client.get(done["download_url"]).status_code == 409
    assert api_client.get("/api/transfers/" + raw["id"]).json()["status"] == "invalidated"
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM transfer_blobs").fetchone()["n"] == 0
    assert not transfer_worker.run_one()


def test_export_rejects_changed_base_and_commits_invalidation(database, api_client):
    lead_id = create_lead(database)
    approve(
        database,
        proposal_for(database, lead_id, "website", "https://nova-empresa.invalid"),
        allow_replace=True,
    )
    queued = api_client.post("/api/exports").json()
    done = drain(api_client, queued["id"])
    with database() as conn:
        service.suppress(conn, lead_id, "teste")
    assert api_client.get(done["download_url"]).status_code == 409
    assert api_client.get("/api/transfers/" + queued["id"]).json()["status"] == "invalidated"
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM transfer_blobs").fetchone()["n"] == 0


def test_hubspot_pages_resume_without_inline_fetch(database, api_client, monkeypatch):
    from radar import hubspot

    monkeypatch.setenv("HUBSPOT_ACCESS_TOKEN", "test-placeholder")
    calls = []

    def list_contacts(self, after=None):
        calls.append(after)
        return {
            "results": [
                {
                    "id": "fake-" + str(after),
                    "properties": {
                        "firstname": "Maria",
                        "lastname": "Exemplo",
                        "company": "Fictícia",
                    },
                }
            ],
            **({"paging": {"next": {"after": "2"}}} if after is None else {}),
        }

    monkeypatch.setattr(hubspot.HubSpotClient, "list_contacts", list_contacts)
    response = api_client.post("/api/hubspot/import", json={})
    assert response.status_code == 202 and calls == []
    done = drain(api_client, response.json()["id"])
    assert calls == [None, "2"] and done["created_count"] == 2 and done["total"] == 2
    with database() as conn:
        assert conn.execute("SELECT count(*) AS n FROM leads").fetchone()["n"] == 2


def test_aggregates_include_zero_and_false_but_exclude_suppressed(database):
    from psycopg.types.json import Jsonb

    from radar.validation import CORE_FIELDS

    lead_id = create_lead(database)
    with database() as conn:
        conn.execute(
            "UPDATE leads SET data=%s,baseline=%s WHERE id=%s",
            (
                Jsonb({"employee_count": 0, "website": "https://ficticia.invalid"}),
                Jsonb({"website": "https://ficticia.invalid"}),
                lead_id,
            ),
        )
        actual = base_queries.coverage(conn)
        assert actual["total"] == 1
        assert (
            next(item for item in actual["fields"] if item["field"] == "employee_count")["after"]
            == 1
        )
        assert actual["completeness"] == round(200 / len(CORE_FIELDS))
        service.suppress(conn, lead_id, "teste")
        assert base_queries.coverage(conn)["total"] == 0


def test_transient_dns_retries_without_holding_privacy_lock(database, api_client, monkeypatch):
    from radar import data_standard

    original = data_standard.prepare_email
    seen = []

    def unavailable(email, cfg):
        with database() as conn:
            assert conn.execute("SELECT pg_try_advisory_xact_lock(907202612) AS free").fetchone()[
                "free"
            ]
        seen.append(email)
        raise data_standard.VerificationUnavailable("Temporário")

    monkeypatch.setattr(data_standard, "prepare_email", unavailable)
    job = upload(api_client, mapping={"Nome": "name", "Empresa": "company"})
    transfer_worker.run_one()
    transfer_worker.run_one()
    result = api_client.get("/api/transfers/" + job["id"]).json()
    assert result["status"] == "queued" and result["processed"] == 0 and result["attempts"] == 1
    monkeypatch.setattr(data_standard, "prepare_email", original)
    with database() as conn:
        conn.execute("UPDATE transfer_jobs SET available_at=now() WHERE id=%s", (job["id"],))
    assert drain(api_client, job["id"])["created_count"] == 205


def test_expiry_cleans_staging_and_completed_export(database, api_client):
    imported = upload(api_client)
    drain(api_client, imported["id"], "awaiting_mapping")
    exported = api_client.post("/api/exports").json()
    drain(api_client, exported["id"])
    with database() as conn:
        conn.execute("UPDATE transfer_jobs SET expires_at=now()-interval '1 second'")
        transfers.cleanup(conn)
        assert conn.execute("SELECT count(*) AS n FROM transfer_blobs").fetchone()["n"] == 0
        assert conn.execute("SELECT count(*) AS n FROM transfer_rows").fetchone()["n"] == 0
    assert (
        api_client.post(
            "/api/transfers/" + imported["id"] + "/start", json={"mapping": {"Nome": "name"}}
        ).status_code
        == 409
    )
    assert api_client.get("/api/transfers/" + exported["id"] + "/download").status_code == 409


def test_numeric_hubspot_id_in_xlsx_is_text_and_reconciles(database, api_client):
    import json

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(["Nome", "ID HubSpot"])
    sheet.append(["Maria Exemplo", 123456789])
    file = io.BytesIO()
    workbook.save(file)
    for _ in range(2):
        response = api_client.post(
            "/api/import",
            files={"file": ("ids.xlsx", file.getvalue())},
            data={"mapping": json.dumps({"Nome": "name", "ID HubSpot": "hubspot_id"})},
        )
        assert response.status_code == 202
        assert drain(api_client, response.json()["id"])["error_count"] == 0
    with database() as conn:
        assert conn.execute("SELECT name,hubspot_id FROM leads").fetchall() == [
            {"name": "Maria Exemplo", "hubspot_id": "123456789"}
        ]


def test_hubspot_rate_limit_uses_retry_queue(database, api_client, monkeypatch):
    import httpx

    from radar import hubspot

    monkeypatch.setenv("HUBSPOT_ACCESS_TOKEN", "test-placeholder")
    calls = []

    def transport(method, url, **kwargs):
        calls.append((method, url))
        return httpx.Response(429 if len(calls) == 1 else 200, json={"results": []})

    monkeypatch.setattr(hubspot.httpx, "request", transport)
    job = api_client.post("/api/hubspot/import", json={}).json()
    transfer_worker.run_one()
    retry = api_client.get("/api/transfers/" + job["id"]).json()
    assert retry["status"] == "queued" and retry["attempts"] == 1 and retry["processed"] == 0
    with database() as conn:
        conn.execute("UPDATE transfer_jobs SET available_at=now() WHERE id=%s", (job["id"],))
    assert drain(api_client, job["id"])["total"] == 0
    assert len(calls) == 2 and all(method == "GET" for method, _ in calls)
