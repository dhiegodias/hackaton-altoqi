import csv
import io
import unicodedata
import zipfile

from openpyxl import load_workbook

from radar.service import capture, digest
from radar.validation import FIELDS

ALIASES = {
    "nome": "name",
    "nome completo": "name",
    "empresa": "company",
    "e-mail": "email",
    "email": "email",
    "site": "website",
    "cargo": "role",
    "telefone": "phone",
    "telefone profissional": "phone",
    "phone": "phone",
    "razao social": "legal_name",
    "porte": "company_size",
    "cidade": "city",
    "uf": "state",
    "segmento": "segment",
    "capital social": "capital_social",
    "funcionarios": "employee_count",
    "pessoas": "employee_count",
    "papel na decisao": "decision_role",
    "id hubspot": "hubspot_id",
    "record id": "hubspot_id",
}


def normalize_header(value):
    plain = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(c for c in plain if not unicodedata.combining(c))


def parse_rows(content, filename):
    if len(content) > 3_000_000:
        raise ValueError("Arquivo excede 3 MB.")
    if filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 20_000_000:
                raise ValueError("Planilha descompactada excede o limite de 20 MB.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        iterator = workbook.active.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(iterator, [])]
        rows = []
        for row in iterator:
            if len(rows) >= 1000:
                raise ValueError("Limite de 1.000 linhas por arquivo.")
            if any(isinstance(cell, str) and cell.startswith("=") for cell in row):
                raise ValueError(
                    "A planilha contém fórmulas. Exporte como valores antes de importar."
                )
            rows.append(dict(zip(headers, row, strict=False)))
        workbook.close()
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = reader.fieldnames or []
        rows = []
        for row in reader:
            if len(rows) >= 1000:
                raise ValueError("Limite de 1.000 linhas por arquivo.")
            if None in row:
                raise ValueError("Linhas com mais valores que o cabeçalho.")
            rows.append(row)
    if not headers or len(headers) != len(set(headers)):
        raise ValueError("Cabeçalho vazio ou com colunas duplicadas.")
    return headers, rows


def automatic_mapping(headers):
    known = {*FIELDS, "name", "company", "email", "hubspot_id"}
    result = {}
    for header in headers:
        normal = normalize_header(header)
        mapped = ALIASES.get(normal, normal)
        if mapped in known:
            result[header] = mapped
    return result


def validate_mapping(mapping):
    targets = [value for value in mapping.values() if value]
    if len(targets) != len(set(targets)):
        raise ValueError("Mapeie cada campo apenas uma vez.")
    if not set(targets).intersection({"name", "email", "company"}):
        raise ValueError("Mapeie nome, empresa ou e-mail.")
    if set(targets) - {*FIELDS, "name", "email", "company", "hubspot_id"}:
        raise ValueError("Mapeamento contém campo desconhecido.")


def row_payload(row, mapping):
    mapped = {
        target: row.get(source)
        for source, target in mapping.items()
        if target and row.get(source) not in (None, "")
    }
    if not mapped:
        return None
    payload = {key: str(mapped.get(key, "")) for key in ("name", "email", "company", "hubspot_id")}
    payload["data"] = {key: value for key, value in mapped.items() if key in FIELDS}
    return payload


def import_rows(conn, rows, mapping, batch_key, origin="csv"):
    validate_mapping(mapping)
    result = {"created": 0, "duplicates": 0, "errors": [], "ids": []}
    for index, row in enumerate(rows, start=2):
        payload = row_payload(row, mapping)
        if not payload:
            continue
        try:
            with conn.transaction():
                key = "import:" + digest([batch_key, index, payload, origin])
                conn.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s, 0))", (key,))
                replay = conn.execute("SELECT 1 FROM requests WHERE key=%s", (key,)).fetchone()
                item = capture(conn, payload, key, origin=origin)
                result["duplicates" if replay or item["duplicate"] else "created"] += 1
                result["ids"].append(item["id"])
        except (ValueError, TypeError) as exc:
            result["errors"].append({"row": index, "message": str(exc)})
    return result
