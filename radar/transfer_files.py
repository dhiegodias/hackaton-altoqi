"""Bounded CSV/XLSX iteration over durable PostgreSQL file chunks."""

import csv
import io
import os
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

CHUNK_BYTES = 256 * 1024
MAX_BYTES = int(os.getenv("RADAR_IMPORT_MAX_BYTES", str(100 * 1024 * 1024)))
MAX_ROWS = int(os.getenv("RADAR_IMPORT_MAX_ROWS", "100000"))
MAX_COLUMNS = 64
MAX_CELL = 4096
MAX_INFLATED = 512 * 1024 * 1024
csv.field_size_limit(MAX_CELL)


class BlobReader(io.RawIOBase):
    """Seekable input for ZipFile/TextIOWrapper, retaining at most one chunk."""

    def __init__(self, conn, job):
        self.conn, self.job_id, self.size = conn, job["id"], job["byte_count"]
        self.position = 0
        self.cached_part, self.cached = None, b""

    def readable(self):
        return True

    def seekable(self):
        return True

    def tell(self):
        return self.position

    def seek(self, offset, whence=0):
        value = offset + (self.position if whence == 1 else self.size if whence == 2 else 0)
        if value < 0:
            raise ValueError("Posição inválida no arquivo.")
        self.position = value
        return value

    def readinto(self, buffer):
        value = self.read(len(buffer))
        buffer[: len(value)] = value
        return len(value)

    def read(self, size=-1):
        remaining = max(0, self.size - self.position)
        size = remaining if size < 0 else min(size, remaining)
        output = bytearray()
        while size:
            part, offset = divmod(self.position, CHUNK_BYTES)
            if part != self.cached_part:
                row = self.conn.execute(
                    "SELECT content FROM transfer_blobs WHERE job_id=%s AND kind='input' AND part=%s",
                    (self.job_id, part),
                ).fetchone()
                if not row:
                    raise ValueError("Arquivo indisponível; tarefa cancelada ou expirada.")
                self.cached_part, self.cached = part, bytes(row["content"])
            take = min(size, len(self.cached) - offset)
            if take <= 0:
                raise ValueError("Arquivo incompleto.")
            output.extend(self.cached[offset : offset + take])
            self.position += take
            size -= take
        return bytes(output)


def headers_checked(values):
    headers = [str(v or "").strip() for v in values]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or len(headers) > MAX_COLUMNS or any(not h or len(h) > 150 for h in headers):
        raise ValueError("Use de 1 a 64 colunas, com cabeçalhos preenchidos de até 150 caracteres.")
    if len(headers) != len(set(headers)):
        raise ValueError("Cabeçalho com colunas duplicadas.")
    return headers


def records(file, filename, encoding="utf-8-sig"):
    """Yield header (line 1), then physical row numbers; never build a full list."""
    workbook = None
    text = None
    try:
        if filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(file) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(i.file_size for i in entries) > MAX_INFLATED:
                    raise ValueError("XLSX descompactado excede 512 MB ou 1.000 partes.")
                if any(
                    i.filename.endswith("sharedStrings.xml") and i.file_size > 32 * 1024 * 1024
                    for i in entries
                ):
                    raise ValueError("Tabela de textos do XLSX excede 32 MB. Exporte em CSV.")
            file.seek(0)
            workbook = load_workbook(file, read_only=True, data_only=False, keep_links=False)
            sheet = workbook.active
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ValueError(
                    "Limite de 64 colunas. Remova colunas/formatos excedentes ou use CSV."
                )
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS + 1)
            raw_headers = list(next(iterator, []))
        else:
            text = io.TextIOWrapper(file, encoding=encoding, newline="")
            sample = text.read(4096)
            text.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            iterator = csv.reader(text, dialect)
            raw_headers = next(iterator, [])
        headers = headers_checked(raw_headers)
        yield 1, headers
        count = 0
        for line, values in enumerate(iterator, start=2):
            if not any(v not in (None, "") for v in values):
                continue
            count += 1
            if count > MAX_ROWS:
                raise ValueError(f"Limite de {MAX_ROWS:,} linhas por arquivo.".replace(",", "."))
            if any(v not in (None, "") for v in values[len(headers) :]):
                raise ValueError(f"Linha {line}: valores além do cabeçalho.")
            values = [
                v.isoformat() if isinstance(v, (date, datetime)) else v
                for v in values[: len(headers)]
            ]
            if any(isinstance(v, str) and v.startswith("=") for v in values):
                raise ValueError(
                    f"Linha {line}: fórmulas não são permitidas; exporte como valores."
                )
            if any(len(str(v or "")) > MAX_CELL for v in values):
                raise ValueError(f"Linha {line}: célula excede {MAX_CELL} caracteres.")
            yield line, dict(zip(headers, values, strict=False))
    except (zipfile.BadZipFile, UnicodeError, csv.Error) as exc:
        raise ValueError("Arquivo inválido, texto malformado ou célula acima do limite.") from exc
    finally:
        if workbook:
            workbook.close()
        if text:
            text.detach()
